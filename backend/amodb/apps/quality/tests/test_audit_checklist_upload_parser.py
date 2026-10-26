from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from amodb.apps.quality.audit_checklist_upload_parser import draft_from_upload


def test_form_metadata_and_merged_footer_are_not_assessment_rows():
    document = Document()
    table = document.add_table(rows=0, cols=6)
    for values in [
        ("", "Date of Audit:", "", "", "", ""),
        ("", "OPERATIONS", "S", "U", "N/C", "N/A"),
        ("4.2", "Verify training records\nInclude recurrent training.", "", "", "", ""),
        ("4.3", "Is the date of audit recorded correctly?", "", "", "", ""),
    ]:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    footer = table.add_row().cells
    footer[0].merge(footer[-1]).text = "Auditor’s Assessment & Recommendations"
    output = BytesIO()
    document.save(output)
    draft = draft_from_upload(output.getvalue(), "source.docx")
    assert [item["checklist_ref"] for item in draft["items"]] == ["4.2", "4.3"]
    assert draft["items"][0]["prompt"] == "Verify training records\nInclude recurrent training."


def test_spaced_word_headers_preserve_yes_no_na_without_cover_questions():
    document = Document()
    cover = document.add_table(rows=0, cols=5)
    for cell, value in zip(cover.add_row().cells, ("", "Post holders", "", "", "")):
        cell.text = value
    table = document.add_table(rows=0, cols=5)
    for values in [("", "WORKSHOP", "Y E S", "N O", "N / A"),
                   ("1 . 1", "Is the approval current?", "", "", ""),
                   ("BATTERY WORKSHOP", "", "", "", ""),
                   ("1.2", "Is the test equipment calibrated?", "", "", "")]:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    output = BytesIO()
    document.save(output)
    draft = draft_from_upload(output.getvalue(), "approved.docx")
    assert [item["checklist_ref"] for item in draft["items"]] == ["1.1", "1.2"]
    assert draft["items"][1]["section"] == "WORKSHOP / BATTERY WORKSHOP"
    assert [option["value"] for option in draft["items"][0]["response_options"]] == ["YES", "NO", "N/A"]


def test_supplied_approved_maintenance_checklist_has_189_source_questions():
    source = Path(r"C:\Users\Admin\Downloads\SL-QMS-29 Approved Maintenance Organization REV 01.docx")
    if not source.exists():
        pytest.skip("User-supplied validation document is local to the live runner")
    draft = draft_from_upload(source.read_bytes(), source.name)
    assert draft["question_count"] == 189
    assert draft["section_count"] == 14
    assert all([option["value"] for option in item["response_options"]] == ["YES", "NO", "N/A"] for item in draft["items"])


def test_actual_contractors_upload_retains_only_source_assessment_rows():
    source = Path(__file__).resolve().parents[4] / "uploads/manuals/safarilink/SL_QMS_20/ID-8YK4VWD6/SL-QMS-20_Contractors_Audit_Checklist.docx"
    if not source.exists():
        pytest.skip("Local controlled upload is not checked into this checkout")
    draft = draft_from_upload(source.read_bytes(), source.name)
    assert draft["question_count"] == 172
    assert draft["section_count"] == 39
    assert draft["items"][0]["section"] == "ORGANIZATION/ MANAGEMENT"
    assert draft["items"][-1]["prompt"] == "LIST OF PERSONEL/CREW AND OPERATIONAL STAFF"
    assert all(item["prompt"].lower() != "date of audit:" for item in draft["items"])
    assert all([option["value"] for option in item["response_options"]] == ["S", "U", "N/C", "N/A"] for item in draft["items"])


def _paper_form() -> bytes:
    document = Document()
    table = document.add_table(rows=0, cols=6)
    for values in [
        ("", "ORGANIZATION/ MANAGEMENT", "S", "U", "N/C", "N/A"),
        ("", "Is there a description and organogram?", "", "", "", ""),
        ("", "Are approved key post holders in place?", "", "", "", ""),
        ("", "Remarks:", "", "", "", ""),
        ("", "FLIGHT SAFETY MANAGEMENT", "S", "U", "N/C", "N/A"),
        ("", "Are risk assessments conducted for all identified hazards?", "", "", "", ""),
        ("", "Remarks:", "", "", "", ""),
    ]:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def test_word_form_extraction_keeps_order_sections_and_explicit_source_vocabulary():
    source = _paper_form()
    draft = draft_from_upload(source, "SL-QMS-20.docx")
    assert draft["question_count"] == 3
    assert draft["section_count"] == 2
    assert [item["checklist_ref"] for item in draft["items"]] == ["1", "2", "3"]
    assert [item["section"] for item in draft["items"]] == [
        "ORGANIZATION/ MANAGEMENT", "ORGANIZATION/ MANAGEMENT", "FLIGHT SAFETY MANAGEMENT",
    ]
    assert draft["items"][2]["prompt"].startswith("Are risk assessments")
    assert [option["value"] for option in draft["items"][0]["response_options"]] == ["S", "U", "N/C", "N/A"]
    assert draft["items"][0]["response_options"][2]["canonical_status"] == "NOT_VERIFIED"
    assert draft["sha256"] and len(draft["sha256"]) == 64


@pytest.mark.parametrize("name,data", [
    ("fake.pdf", b"not a PDF"),
    ("fake.docx", b"not a Word file"),
    ("broken.docx", b"PKbroken archive"),
    ("broken.pdf", b"%PDFbroken file"),
    ("bad.txt", b"anything"),
    ("empty.docx", b""),
])
def test_invalid_uploads_fail_closed(name: str, data: bytes):
    with pytest.raises(ValueError):
        draft_from_upload(data, name)


def _pdf_form() -> bytes:
    from reportlab.pdfgen import canvas

    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=(595, 842))
    for index, line in enumerate([
        "ORGANIZATION/ MANAGEMENT S U N/C N/A",
        "1. Is there a description and organogram?",
        "2. Are approved key post holders in place with",
        "defined responsibilities and authorization?",
        "Remarks:",
        "FLIGHT SAFETY MANAGEMENT S U N/C N/A",
        "3. Are risk assessments conducted for all identified",
        "hazards and organizational changes?",
        "Remarks:",
        "RECURRENT TRAINING/ SIX MONTHLY PROFICIENCY",
        "CHECKS",
        "S U N/C N/A",
        "4. Are training records reviewed and signed?",
    ]):
        pdf.drawString(35, 800 - 22 * index, line)
    pdf.save()
    return output.getvalue()


def test_pdf_import_preserves_multiline_questions_and_wrapped_section_headers():
    draft = draft_from_upload(_pdf_form(), "Audit checklist.pdf")
    assert [item["checklist_ref"] for item in draft["items"]] == ["1", "2", "3", "4"]
    assert draft["question_count"] == 4
    assert draft["section_count"] == 3
    assert draft["items"][1]["prompt"].endswith("defined responsibilities and authorization?")
    assert draft["items"][2]["section"] == "FLIGHT SAFETY MANAGEMENT"
    assert draft["items"][2]["prompt"].endswith("hazards and organizational changes?")
    assert draft["items"][3]["section"] == "RECURRENT TRAINING/ SIX MONTHLY PROFICIENCY CHECKS"
    assert draft["items"][0]["response_options"][0]["value"] == "S"
