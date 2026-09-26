from __future__ import annotations

import html
import io
import os
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from xml.etree import ElementTree

_MAX_INDEX_CHARS = int(os.getenv("DOCUMENT_TEXT_INDEX_MAX_CHARS", "2000000"))
_TIKA_URL = os.getenv("DOCUMENT_TIKA_URL", "").rstrip("/")
_TIKA_TIMEOUT = float(os.getenv("DOCUMENT_TIKA_TIMEOUT_SECONDS", "15"))


@dataclass(frozen=True)
class ExtractedDocumentText:
    text: str
    engine: str
    truncated: bool
    warning: str | None = None


def _bounded(value: str, engine: str, warning: str | None = None) -> ExtractedDocumentText:
    compact = re.sub(r"[\t\r ]+", " ", value or "")
    compact = re.sub(r"\n{3,}", "\n\n", compact).strip()
    truncated = len(compact) > _MAX_INDEX_CHARS
    return ExtractedDocumentText(
        text=compact[:_MAX_INDEX_CHARS],
        engine=engine,
        truncated=truncated,
        warning=warning,
    )


def _tika_extract(content: bytes, mime_type: str) -> ExtractedDocumentText | None:
    """Use an optional Apache Tika sidecar without making it a runtime dependency.

    Tika remains a derived-index service: the original tenant file stays in the
    DMS and can always be re-indexed if the extractor is replaced or upgraded.
    """
    if not _TIKA_URL:
        return None
    request = Request(
        f"{_TIKA_URL}/tika",
        data=content,
        headers={
            "Accept": "text/plain",
            "Content-Type": mime_type or "application/octet-stream",
            "User-Agent": "AMO-Portal-DMS/1.0",
        },
        method="PUT",
    )
    try:
        with urlopen(request, timeout=_TIKA_TIMEOUT) as response:
            if int(getattr(response, "status", 200)) >= 300:
                return None
            payload = response.read(max(4_000_000, _MAX_INDEX_CHARS * 4)).decode("utf-8", errors="replace")
            return _bounded(payload, "APACHE_TIKA")
    except (HTTPError, URLError, TimeoutError, OSError):
        return None


def _xml_archive_text(content: bytes, prefixes: tuple[str, ...]) -> str:
    values: list[str] = []
    length = 0
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for name in archive.namelist():
            if not any(name.startswith(prefix) for prefix in prefixes) or not name.endswith(".xml"):
                continue
            try:
                root = ElementTree.fromstring(archive.read(name))
            except ElementTree.ParseError:
                continue
            for node in root.iter():
                if node.text and node.text.strip():
                    values.append(node.text.strip())
                    length += len(node.text.strip())
                    if length >= _MAX_INDEX_CHARS:
                        return "\n".join(values)
    return "\n".join(values)


def _pdf_text(content: bytes) -> str:
    try:
        import fitz  # type: ignore
    except ImportError:
        return ""
    values: list[str] = []
    length = 0
    with fitz.open(stream=content, filetype="pdf") as document:
        for page in document:
            values.append(str(page.get_text("text") or ""))
            length += len(values[-1])
            if length >= _MAX_INDEX_CHARS:
                break
    return "\n".join(values)


def _xlsx_text(content: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    values: list[str] = []
    length = 0
    try:
        for sheet in workbook.worksheets:
            values.append(sheet.title)
            for row in sheet.iter_rows(values_only=True):
                line = " | ".join(str(value) for value in row if value is not None)
                if line:
                    values.append(line)
                    length += len(line)
                if length >= _MAX_INDEX_CHARS:
                    return "\n".join(values)
    finally:
        workbook.close()
    return "\n".join(values)


def extract_document_text(filename: str, content: bytes, mime_type: str) -> ExtractedDocumentText:
    """Extract searchable text locally, with optional Tika as the broad-format engine."""
    tika = _tika_extract(content, mime_type)
    if tika and tika.text:
        return tika

    suffix = os.path.splitext(filename.lower())[1]
    try:
        if suffix == ".pdf" or mime_type == "application/pdf":
            value = _pdf_text(content)
            return _bounded(value, "PYMUPDF", None if value.strip() else "PDF has no embedded searchable text; OCR indexing is required.")
        if suffix == ".docx":
            return _bounded(_xml_archive_text(content, ("word/",)), "OOXML_DOCX")
        if suffix == ".pptx":
            return _bounded(_xml_archive_text(content, ("ppt/slides/", "ppt/notesSlides/")), "OOXML_PPTX")
        if suffix in {".xlsx", ".xlsm"}:
            return _bounded(_xlsx_text(content), "OPENPYXL")
        if suffix in {".txt", ".csv", ".md", ".json", ".xml", ".yaml", ".yml", ".log"} or mime_type.startswith("text/"):
            return _bounded(content.decode("utf-8", errors="replace"), "TEXT")
        if suffix in {".html", ".htm"}:
            decoded = content.decode("utf-8", errors="replace")
            decoded = re.sub(r"<script\b[^>]*>.*?</script>", " ", decoded, flags=re.IGNORECASE | re.DOTALL)
            decoded = re.sub(r"<style\b[^>]*>.*?</style>", " ", decoded, flags=re.IGNORECASE | re.DOTALL)
            return _bounded(html.unescape(re.sub(r"<[^>]+>", " ", decoded)), "HTML")
    except (OSError, ValueError, zipfile.BadZipFile):
        return _bounded("", "LOCAL_FALLBACK", "The file could not be parsed for full-text indexing.")

    return _bounded("", "METADATA_ONLY", "No local full-text extractor is configured for this file type.")
